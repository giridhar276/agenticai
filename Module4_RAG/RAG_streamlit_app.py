import hashlib
from io import BytesIO

import streamlit as st
from docx import Document as WordDocument
from pypdf import PdfReader
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

st.set_page_config(page_title="Document RAG", page_icon="📚", layout="wide")
st.title("📚 Ask your documents")
st.write("Upload PDF, DOCX, TXT, or Markdown files, then ask questions using their content.")

with st.sidebar:
    st.header("Settings")
    api_key = st.text_input("OpenAI API key", type="password", help="Used for embeddings and answers. It is kept in this Streamlit session only.")
    chat_model = st.text_input("Chat model", value="gpt-4.1-mini")
    embedding_model = st.text_input("Embedding model", value="text-embedding-3-small")
    chunk_size = st.slider("Chunk size", min_value=300, max_value=1500, value=800, step=100)
    chunk_overlap = st.slider("Chunk overlap", min_value=0, max_value=300, value=100, step=25)
    top_k = st.slider("Retrieved chunks", min_value=1, max_value=8, value=4)

uploads = st.file_uploader(
    "Choose one or more documents",
    type=["pdf", "docx", "txt", "md", "markdown"],
    accept_multiple_files=True,
)


def upload_signature(files):
    """Identify the current upload set so stale indexes are not used."""
    digest = hashlib.sha256()
    for file in files:
        digest.update(file.name.encode("utf-8"))
        digest.update(file.getvalue())
    return digest.hexdigest()


def read_upload(file):
    """Return one LangChain Document per page or text file."""
    suffix = file.name.rsplit(".", 1)[-1].lower()
    data = file.getvalue()
    if suffix == "pdf":
        reader = PdfReader(BytesIO(data))
        return [
            Document(page_content=page.extract_text() or "", metadata={"source": file.name, "page": i})
            for i, page in enumerate(reader.pages, start=1)
        ]
    if suffix == "docx":
        word_doc = WordDocument(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in word_doc.paragraphs if paragraph.text.strip())
        return [Document(page_content=text, metadata={"source": file.name})]
    text = data.decode("utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": file.name})]


current_signature = upload_signature(uploads) if uploads else None
if st.session_state.get("indexed_signature") and current_signature != st.session_state.indexed_signature:
    st.info("Your uploaded files changed. Build the index again before asking questions.")

build_disabled = not uploads or not api_key.strip()
if st.button("Build / rebuild document index", type="primary", disabled=build_disabled):
    try:
        with st.spinner("Reading documents and creating the search index…"):
            documents = []
            for uploaded_file in uploads:
                documents.extend(read_upload(uploaded_file))
            documents = [doc for doc in documents if doc.page_content.strip()]
            if not documents:
                raise ValueError("No readable text was found in the uploaded files.")

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=chunk_size,
                chunk_overlap=min(chunk_overlap, chunk_size - 1),
            )
            chunks = splitter.split_documents(documents)
            embeddings = OpenAIEmbeddings(model=embedding_model, api_key=api_key)
            vector_db = Chroma.from_documents(documents=chunks, embedding=embeddings)
            llm = ChatOpenAI(model=chat_model, temperature=0, api_key=api_key)
            prompt = ChatPromptTemplate.from_messages([
                ("system", "Answer the question using only the context below. If the context does not contain the answer, say you could not find it in the uploaded documents. Be concise.\n\nContext:\n{context}"),
                ("human", "{question}"),
            ])
            chain = prompt | llm | StrOutputParser()

            st.session_state.vector_db = vector_db
            st.session_state.rag_chain = chain
            st.session_state.indexed_signature = current_signature
            st.session_state.top_k = top_k
            st.session_state.messages = []
        st.success(f"Indexed {len(uploads)} file(s) into {len(chunks)} searchable chunks.")
    except Exception as exc:
        st.error(f"Could not build the index: {exc}")

if st.session_state.get("rag_chain") and current_signature == st.session_state.get("indexed_signature"):
    st.divider()
    st.subheader("Chat with your documents")
    for message in st.session_state.get("messages", []):
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
            if message.get("sources"):
                with st.expander("Sources"):
                    for source in message["sources"]:
                        st.write(source)

    question = st.chat_input("Ask a question about the uploaded files")
    if question:
        st.session_state.messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        with st.chat_message("assistant"):
            with st.spinner("Searching documents…"):
                try:
                    matches = st.session_state.vector_db.similarity_search(question, k=top_k)
                    context = "\n\n".join(doc.page_content for doc in matches)
                    answer = st.session_state.rag_chain.invoke({"question": question, "context": context})
                    sources = sorted({
                        f"{doc.metadata.get('source', 'Unknown file')}"
                        + (f", page {doc.metadata['page']}" if doc.metadata.get("page") else "")
                        for doc in matches
                    })
                    st.markdown(answer)
                    with st.expander("Sources"):
                        for source in sources:
                            st.write(source)
                    st.session_state.messages.append({"role": "assistant", "content": answer, "sources": sources})
                except Exception as exc:
                    st.error(f"Could not answer the question: {exc}")
else:
    st.caption("Add an OpenAI API key, upload documents, and build the index to start chatting.")
